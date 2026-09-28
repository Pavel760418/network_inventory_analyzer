# -*- coding: utf-8 -*-
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from src.network_inventory_analyzer.coverage import completeness_row
from src.network_inventory_analyzer.excel_exporter import assert_no_excel_errors, export_workbook


def test_category_reserve_reconciles(bundle):
    reserve = bundle["reserve"]
    assert abs(float(reserve["Резерв, ₽"].sum()) - bundle["kpi"]["reserve"]) < 0.02
    assert reserve["Метод расчёта"].str.len().gt(10).all()
    assert reserve["Источник"].ne("").all()
    assert reserve["Статус"].eq("рассчитано").all()


def test_rating_components_and_household_exclusion(bundle):
    rating = bundle["rating"]
    assert not rating.empty
    delta = (rating["сумма_компонент"] - rating["store_score"]).abs()
    assert delta.max() < 0.2
    stores = set(rating["магазин"])
    assert "Магазин Альфа" in stores
    household_names = set(bundle["household"]["наименование"])
    assert "Пакет майка" in household_names
    rated_names = set(bundle["lines"].loc[bundle["lines"]["include_in_store_rating"] == True, "наименование"])  # noqa: E712
    assert "Пакет майка" not in rated_names
    assert "Контейнер пластиковый" not in rated_names


def test_beef_confirmed_only(bundle):
    assert float(bundle["beef_shortages"]["недостача_сумма"].sum()) == 2000.0
    review = bundle["beef_full"]
    broth = review[review["наименование"] == "Бульон говяжий demo"]
    assert not broth.empty
    assert broth.iloc[0]["beef_confirmed"] == False  # noqa: E712
    full_shortage = bundle["lines"].loc[
        (bundle["lines"]["наименование"] == "Говядина вырезка demo") & (bundle["lines"]["недостача_сумма"] > 0),
        "недостача_сумма",
    ].sum()
    assert float(full_shortage) == float(bundle["beef_shortages"]["недостача_сумма"].sum())


def test_posted_without_file_is_explicit(bundle):
    assert bundle["posted_full"].empty
    assert "нет файла" in str(bundle["posted_summary"].iloc[0]["группа"])


def test_completeness_marks_equal_sets(bundle):
    bad = bundle["completeness"][bundle["completeness"]["Статус"] != "OK"]
    assert bad.empty
    row = completeness_row("пример", 3, 3)
    assert row["Статус"] == "OK"
    clipped = completeness_row("пример", 3, 1)
    assert clipped["Статус"] == "ОБРЕЗАНО"


def test_excel_contains_full_sheets(bundle):
    raw = export_workbook(bundle)
    assert assert_no_excel_errors(raw) == []
    wb = load_workbook(BytesIO(raw))
    required = [
        "Резюме_для_руководителя",
        "Дашборд_ключевые_риски",
        "Недостачи_полный_список",
        "Излишки_полный_список",
        "Пересорты_полный_список",
        "Перекрытия_полный_список",
        "Кандидаты_на_пересорт",
        "Хозтовары_упаковка_тара",
        "Товары_без_классификации",
        "Резерв_категорий_детализация",
        "Методология_рейтинга_магазинов",
        "Рейтинг_магазинов_детализация",
        "Аномалии_цен_полный_список",
        "Говядина_Резюме",
        "Говядина_Недостачи",
        "Говядина_Полный_список",
        "Говядина_Методология",
        "Полнота_выгрузки_и_контроль",
        "Методология",
        "Лог_проверок",
        "Оприход_излишки_Резюме",
        "Оприход_излишки_Полный",
    ]
    for name in required:
        assert name in wb.sheetnames
    shortages = wb["Недостачи_полный_список"]
    data_rows = shortages.max_row - 4
    assert data_rows == len(bundle["shortages"])
    assert data_rows > 10 or data_rows == len(bundle["shortages"])
