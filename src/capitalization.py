# -*- coding: utf-8 -*-
"""Parse and match shift-close surplus capitalization (оприходование излишков).

Rules (Release 4):
- exact name matching only (strip + normalized sku_key equality);
- never mix warehouses / stores;
- warehouse→store mapping must be unique; ambiguous → uncertain.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.text_normalize import norm_txt, sku_key

STATUS_MATCHED = "matched"
STATUS_MATCHED_SHORTAGE = "matched_with_shortage"
STATUS_UNMATCHED_NAME = "unmatched_name"
STATUS_UNMATCHED_STORE = "unmatched_store"
STATUS_UNCERTAIN_STORE = "uncertain_store"

_WAREHOUSE_PREFIXES = ("склад", "рц", "фабрика")
_STOP_TOKENS = {
    "склад", "компании", "магазин", "производство", "зеленого", "яблока",
    "и", "в", "на", "от", "по", "г", "ул",
}


def _to_float(v) -> float:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _is_warehouse_header(name: str) -> bool:
    low = name.lower().strip()
    return any(low.startswith(p) for p in _WAREHOUSE_PREFIXES)


def parse_capitalization_excel(filepath: str, sheet_name: Optional[str] = None) -> pd.DataFrame:
    """Parse TDSheet-like capitalization export into flat rows per warehouse+SKU.

    Expected columns (row with headers):
    - Склад компании / Номенклатура in col A
    - Закрытие смены количество
    - Закрытие смены сумма
    """
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(
            "Файл оприходования не найден.\n"
            f"Ожидался файл: {path}\n"
            "Выберите файл оприходования из папки исходных Excel "
            "или проверьте, что он не был перемещён."
        )

    wb = load_workbook(filepath, read_only=False, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    levels = {
        int(idx): int(getattr(dim, "outlineLevel", 0) or 0)
        for idx, dim in ws.row_dimensions.items()
    }

    # locate header row
    header_row = None
    qty_col = 4
    sum_col = 6
    for i in range(1, min(30, ws.max_row + 1)):
        vals = [str(ws.cell(i, c).value or "").lower() for c in range(1, 9)]
        joined = " | ".join(vals)
        if "закрытие смены количество" in joined and "закрытие смены сумма" in joined:
            header_row = i
            for c in range(1, 9):
                v = str(ws.cell(i, c).value or "").lower()
                if "закрытие смены количество" in v:
                    qty_col = c
                if "закрытие смены сумма" in v:
                    sum_col = c
            break
    if header_row is None:
        wb.close()
        raise ValueError(
            "Файл оприходования не распознан: нет колонок "
            "«Закрытие смены количество» / «Закрытие смены сумма»."
        )

    rows: List[dict] = []
    current_wh = ""
    for i in range(header_row + 1, ws.max_row + 1):
        name_raw = ws.cell(i, 1).value
        if name_raw is None or str(name_raw).strip() == "":
            continue
        name = str(name_raw).strip()
        if name.lower() == "номенклатура":
            continue
        lv = levels.get(i, 0)
        qty = _to_float(ws.cell(i, qty_col).value)
        sm = _to_float(ws.cell(i, sum_col).value)

        if lv == 0 and _is_warehouse_header(name):
            current_wh = name
            continue
        if _is_warehouse_header(name) and lv == 0:
            current_wh = name
            continue
        # product under warehouse
        if not current_wh:
            # orphan product without warehouse context
            rows.append({
                "склад": "",
                "наименование": name,
                "оп_количество": qty,
                "оп_сумма": sm,
                "sku_key": sku_key(name),
                "name_exact": name,
            })
            continue
        if lv >= 1 or not _is_warehouse_header(name):
            if _is_warehouse_header(name):
                current_wh = name
                continue
            rows.append({
                "склад": current_wh,
                "наименование": name,
                "оп_количество": qty,
                "оп_сумма": sm,
                "sku_key": sku_key(name),
                "name_exact": name,
            })

    wb.close()
    if not rows:
        raise ValueError("В файле оприходования не найдено товарных строк.")
    return pd.DataFrame(rows)


def _significant_tokens(text: str) -> Set[str]:
    n = norm_txt(text)
    toks = set(re.findall(r"[a-zа-я0-9]+", n))
    return {t for t in toks if t not in _STOP_TOKENS and (len(t) > 2 or t.isdigit())}


def map_warehouse_to_store(warehouse: str, inventory_stores: Iterable[str]) -> Tuple[Optional[str], str]:
    """Map capitalization warehouse to a single inventory store.

    Returns (store_name_or_None, mapping_status).
    mapping_status: mapped | unmatched_store | uncertain_store
    """
    stores = [str(s).strip() for s in inventory_stores if str(s).strip()]
    if not warehouse or not stores:
        return None, STATUS_UNMATCHED_STORE

    wh = warehouse.strip()
    wh_low = wh.lower()

    # 1) exact (case-insensitive)
    exact = [s for s in stores if s.lower() == wh_low]
    if len(exact) == 1:
        return exact[0], "mapped"
    if len(exact) > 1:
        return None, STATUS_UNCERTAIN_STORE

    # 2) containment either way
    contained = [
        s for s in stores
        if s.lower() in wh_low or wh_low in s.lower()
    ]
    if len(contained) == 1:
        return contained[0], "mapped"
    if len(contained) > 1:
        return None, STATUS_UNCERTAIN_STORE

    # 3) token overlap (numbers / distinctive words)
    wh_toks = _significant_tokens(wh)
    scored: List[Tuple[int, str]] = []
    for s in stores:
        st = _significant_tokens(s)
        inter = wh_toks & st
        if not inter:
            continue
        # Prefer digit matches (store numbers)
        score = sum(10 if t.isdigit() else 3 for t in inter)
        scored.append((score, s))
    scored.sort(key=lambda x: (-x[0], x[1]))
    if not scored:
        return None, STATUS_UNMATCHED_STORE
    best_score = scored[0][0]
    top = [s for sc, s in scored if sc == best_score]
    if len(top) == 1 and best_score >= 10:
        return top[0], "mapped"
    if len(top) == 1 and best_score >= 3:
        # weak textual match — still unique
        return top[0], "mapped"
    return None, STATUS_UNCERTAIN_STORE


def _build_inventory_indexes(inv_df: pd.DataFrame) -> Dict[str, Dict[str, pd.DataFrame]]:
    """store -> { 'by_exact': dict name->df, 'by_key': dict sku_key->df } aggregated views."""
    indexes: Dict[str, Dict[str, pd.DataFrame]] = {}
    if inv_df.empty:
        return indexes
    work = inv_df.copy()
    if "sku_key" not in work.columns:
        work["sku_key"] = work["наименование"].map(sku_key)
    work["name_exact"] = work["наименование"].astype(str).str.strip()
    for store, grp in work.groupby("магазин"):
        indexes[str(store)] = {
            "by_exact": {n: g for n, g in grp.groupby("name_exact", dropna=False)},
            "by_key": {k: g for k, g in grp.groupby("sku_key", dropna=False) if k},
        }
    return indexes


def match_capitalization_to_inventory(
    cap_df: pd.DataFrame,
    inventory_df: pd.DataFrame,
) -> pd.DataFrame:
    """Exact-name match within mapped store only. Never cross stores."""
    stores = sorted(inventory_df["магазин"].dropna().astype(str).unique()) if not inventory_df.empty else []
    indexes = _build_inventory_indexes(inventory_df)

    # cache warehouse mapping
    wh_map: Dict[str, Tuple[Optional[str], str]] = {}
    out_rows = []

    for _, row in cap_df.iterrows():
        wh = str(row.get("склад", "") or "").strip()
        name = str(row.get("наименование", "") or "").strip()
        qty = float(row.get("оп_количество", 0) or 0)
        sm = float(row.get("оп_сумма", 0) or 0)
        key = str(row.get("sku_key") or sku_key(name))

        if wh not in wh_map:
            wh_map[wh] = map_warehouse_to_store(wh, stores)
        store, map_status = wh_map[wh]

        comment = ""
        linked_shortage = False
        inv_surplus = 0.0
        inv_shortage = 0.0
        match_docs = ""

        if map_status == STATUS_UNMATCHED_STORE:
            status = STATUS_UNMATCHED_STORE
            comment = "Склад оприходования не сопоставлен ни с одним магазином инвентаризации"
        elif map_status == STATUS_UNCERTAIN_STORE:
            status = STATUS_UNCERTAIN_STORE
            comment = "Неоднозначное сопоставление склада — данные не смешивались"
        else:
            assert store is not None
            idx = indexes.get(store, {})
            hit = idx.get("by_exact", {}).get(name)
            if hit is None or hit.empty:
                hit = idx.get("by_key", {}).get(key)
            if hit is None or hit.empty:
                status = STATUS_UNMATCHED_NAME
                comment = f"Наименование не найдено в инвентаризации склада «{store}» (точное совпадение)"
            else:
                inv_surplus = float(hit["излишек_сумма"].sum()) if "излишек_сумма" in hit.columns else 0.0
                inv_shortage = float(hit["недостача_сумма"].sum()) if "недостача_сумма" in hit.columns else 0.0
                shortage_qty = float(hit["недостача_кол"].sum()) if "недостача_кол" in hit.columns else 0.0
                linked_shortage = inv_shortage > 0 or shortage_qty > 0
                if "документ" in hit.columns:
                    match_docs = "; ".join(sorted(set(hit["документ"].astype(str).head(5))))
                if linked_shortage:
                    status = STATUS_MATCHED_SHORTAGE
                    comment = "Сопоставлено; по этой позиции в инвентаризации есть недостача"
                else:
                    status = STATUS_MATCHED
                    comment = "Сопоставлено по точному наименованию внутри склада"

        out_rows.append({
            "склад": wh,
            "магазин": store or "",
            "store_map_status": map_status if map_status != "mapped" else "mapped",
            "наименование": name,
            "оп_количество": qty,
            "оп_сумма": sm,
            "статус": status,
            "связь_с_недостачей": "ДА" if linked_shortage else "НЕТ",
            "инв_излишек_сумма": inv_surplus,
            "инв_недостача_сумма": inv_shortage,
            "документы": match_docs,
            "комментарий": comment,
            "sku_key": key,
        })

    result = pd.DataFrame(out_rows)
    if result.empty:
        return result
    return result.sort_values(["склад", "статус", "оп_сумма"], ascending=[True, True, False])


def capitalization_summary(matched_df: pd.DataFrame) -> Dict[str, float]:
    if matched_df is None or matched_df.empty:
        return {
            "cap_rows": 0,
            "cap_qty": 0.0,
            "cap_sum": 0.0,
            "matched": 0,
            "matched_shortage": 0,
            "unmatched_name": 0,
            "unmatched_store": 0,
            "uncertain_store": 0,
            "warehouses": 0,
            "matched_sum": 0.0,
            "shortage_linked_sum": 0.0,
        }
    st = matched_df["статус"]
    return {
        "cap_rows": int(len(matched_df)),
        "cap_qty": float(matched_df["оп_количество"].sum()),
        "cap_sum": float(matched_df["оп_сумма"].sum()),
        "matched": int((st == STATUS_MATCHED).sum()),
        "matched_shortage": int((st == STATUS_MATCHED_SHORTAGE).sum()),
        "unmatched_name": int((st == STATUS_UNMATCHED_NAME).sum()),
        "unmatched_store": int((st == STATUS_UNMATCHED_STORE).sum()),
        "uncertain_store": int((st == STATUS_UNCERTAIN_STORE).sum()),
        "warehouses": int(matched_df["склад"].nunique()),
        "matched_sum": float(
            matched_df.loc[st.isin([STATUS_MATCHED, STATUS_MATCHED_SHORTAGE]), "оп_сумма"].sum()
        ),
        "shortage_linked_sum": float(
            matched_df.loc[st == STATUS_MATCHED_SHORTAGE, "оп_сумма"].sum()
        ),
    }


def enrich_store_metrics_with_capitalization(
    store_metrics: pd.DataFrame,
    matched_df: pd.DataFrame,
) -> pd.DataFrame:
    """Attach capitalization totals to store ranking (by mapped inventory store)."""
    sm = store_metrics.copy()
    if sm.empty:
        return sm
    sm["оприходование_сумма"] = 0.0
    sm["оприходование_кол"] = 0.0
    sm["оприходование_строк"] = 0
    sm["оп_связь_с_недостачей"] = 0
    if matched_df is None or matched_df.empty:
        return sm

    usable = matched_df[matched_df["магазин"].astype(str).str.len() > 0]
    if usable.empty:
        return sm
    agg = usable.groupby("магазин").agg(
        оприходование_сумма=("оп_сумма", "sum"),
        оприходование_кол=("оп_количество", "sum"),
        оприходование_строк=("наименование", "count"),
        оп_связь_с_недостачей=("связь_с_недостачей", lambda s: int((s == "ДА").sum())),
    )
    sm = sm.set_index("магазин")
    for col in agg.columns:
        sm[col] = agg[col]
    sm[list(agg.columns)] = sm[list(agg.columns)].fillna(0)
    return sm.reset_index()
